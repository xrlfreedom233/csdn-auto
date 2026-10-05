from __future__ import annotations

import unittest
import re
from io import BytesIO
from datetime import datetime, timedelta, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock

import requests

from csdn_auto.ai import (
    AIClient,
    GeneratedArticle,
    article_quality_issues,
    parse_article,
)
from csdn_auto.config import Settings
from csdn_auto.csdn import CSDNClient
from csdn_auto.prompts import BASE_SYSTEM_PROMPT, PromptComposer
from csdn_auto.rag import (
    LocalKnowledgeBase,
    OpenAIEmbeddingProvider,
    PersistentVectorKnowledgeBase,
    format_context,
)
from csdn_auto.scheduler import _is_due, _last_success, _mark_success
from csdn_auto.importer import KnowledgeImporter, clean_docx, clean_html, clean_pdf


def settings() -> Settings:
    return Settings(
        ai_base_url="https://example.test/v1",
        ai_api_key="test",
        ai_model="test",
        ai_connect_timeout=5,
        ai_timeout=10,
        csdn_cookie="SESSION=test",
        csdn_x_ca_key="203803574",
        csdn_x_ca_secret="secret",
        csdn_x_ca_nonce="",
        csdn_x_ca_signature="",
        csdn_article_id="",
        csdn_tags="python",
        csdn_category="",
        prompt_history_path="/tmp/csdn_auto-test-prompt-history.json",
        prompt_scene_count=1,
        prompt_tech_min=2,
        prompt_tech_max=3,
        prompt_recent_window=30,
        schedule_interval_hours=12,
        schedule_timezone="Asia/Shanghai",
        rag_enabled=False,
    )


class ParseArticleTests(unittest.TestCase):
    def test_parses_title_and_description(self) -> None:
        article = parse_article(
            "# Python 定时任务\n\n这是正文内容，用于测试描述生成。",
            tags="Python,定时任务",
        )
        self.assertEqual(article.title, "Python 定时任务")
        self.assertTrue(article.description.startswith("这是正文内容"))
        self.assertEqual(article.tags, "Python,定时任务")

    def test_rejects_missing_title(self) -> None:
        with self.assertRaises(ValueError):
            parse_article("没有一级标题")

    def test_quality_gate_accepts_interview_article(self) -> None:
        article = """# MCP 面试高频考点：Tool 安全设计

这是面试中经常被追问的一道题。

**面试官**：如何设计？

**候选人**：先做服务端授权。

## 面试官点评

- 考察点：安全边界
- 合格回答：能说明授权
- 加分项：能说明审计
"""
        self.assertEqual(article_quality_issues(article), [])

    def test_quality_gate_rejects_hallucination_and_internal_paths(self) -> None:
        article = """# MCP 技术文

使用 Java MCP SDK 0.6.1 可减少 90% 的工作，资料位于 imported/Tools-abcd1234.md。
"""
        issues = article_quality_issues(article)
        self.assertTrue(any("标题" in issue for issue in issues))
        self.assertTrue(any("百分比" in issue for issue in issues))
        self.assertTrue(any("SDK 版本" in issue for issue in issues))
        self.assertTrue(any("内部路径" in issue for issue in issues))

    def test_quality_gate_requires_citations_when_rag_is_used(self) -> None:
        article = """# 大厂 MCP 面试常考：Tool 安全设计

这是面试高频考点。

面试官：如何设计？候选人：需要授权。

考察点：安全。合格回答：授权。加分项：审计。
"""
        issues = article_quality_issues(article, require_citations=True)
        self.assertTrue(any("事实引用" in issue for issue in issues))
        self.assertTrue(any("参考资料" in issue for issue in issues))


class AIClientTests(unittest.TestCase):
    def test_reports_proxy_error(self) -> None:
        session = Mock()
        session.post.side_effect = requests.exceptions.ProxyError("proxy unavailable")
        client = AIClient(settings(), session=session)
        with self.assertRaisesRegex(RuntimeError, "无法连接系统代理"):
            client.generate()


class PromptComposerTests(unittest.TestCase):
    def test_combines_enumerated_items_and_records_json_history(self) -> None:
        with TemporaryDirectory() as directory:
            composer = PromptComposer(Path(directory) / "prompt_history.json")
            selection = composer.choose(1, 2, 3, 30)
            self.assertEqual(len(selection.scenes), 1)
            self.assertGreaterEqual(len(selection.technologies), 2)
            self.assertIn("业务场景", selection.user_prompt())
            composer.record(selection)
            self.assertIn(selection.combination_key, composer._recent_keys(30))
            self.assertTrue((Path(directory) / "prompt_history.json").exists())

    def test_generates_mcp_prompt_with_rag_context(self) -> None:
        with TemporaryDirectory() as directory:
            selection = PromptComposer(Path(directory) / "history.json").choose(
                1, 2, 3, 0
            )
            prompt = selection.user_prompt("[资料1 | 来源: test.md]\nMCP 资料")
            self.assertIn("MCP 面试型技术文章", prompt)
            self.assertIn("检索到的参考片段", prompt)
            self.assertTrue(selection.retrieval_query)
            self.assertIn("面试官", prompt)
            self.assertIn("MCP 面试高频考点", BASE_SYSTEM_PROMPT)
            self.assertNotIn("模拟面试", BASE_SYSTEM_PROMPT)


class LocalKnowledgeBaseTests(unittest.TestCase):
    def test_retrieves_relevant_chunks_and_keeps_source(self) -> None:
        with TemporaryDirectory() as directory:
            knowledge = Path(directory)
            (knowledge / "transport.md").write_text(
                "MCP stdio 传输通过标准输入输出与子进程通信。",
                encoding="utf-8",
            )
            (knowledge / "other.txt").write_text(
                "数据库索引优化的一般介绍。", encoding="utf-8"
            )
            chunks = LocalKnowledgeBase(knowledge, chunk_size=200).retrieve(
                "MCP stdio 传输", top_k=1
            )
            self.assertEqual(len(chunks), 1)
            self.assertEqual(chunks[0].source, "transport.md")
            self.assertIn("标准输入输出", format_context(chunks))

    def test_context_exposes_title_and_public_url_not_internal_path(self) -> None:
        with TemporaryDirectory() as directory:
            knowledge = Path(directory)
            (knowledge / "Tools-abcd1234.md").write_text(
                "<!-- source: https://raw.githubusercontent.com/modelcontextprotocol/"
                "modelcontextprotocol/main/docs/specification/2026-07-28/server/tools.mdx -->\n\n"
                "# Tools\n\nMCP Tools are model-controlled operations with input schemas.",
                encoding="utf-8",
            )
            chunks = LocalKnowledgeBase(knowledge).retrieve("MCP Tools", top_k=1)
            context = format_context(chunks)
            self.assertIn("标题: Tools", context)
            self.assertIn(
                "https://modelcontextprotocol.io/specification/2026-07-28/server/tools",
                context,
            )
            self.assertNotIn("Tools-abcd1234.md", context)
            self.assertNotIn("imported/", context)
            self.assertNotIn("raw.githubusercontent.com", context)

    def test_missing_directory_falls_back_to_empty_results(self) -> None:
        with TemporaryDirectory() as directory:
            missing = Path(directory) / "missing"
            self.assertEqual(LocalKnowledgeBase(missing).retrieve("MCP"), [])


class VectorKnowledgeBaseTests(unittest.TestCase):
    def test_incrementally_indexes_and_retrieves_with_source(self) -> None:
        with TemporaryDirectory() as directory:
            document_path = Path(directory) / "mcp.md"
            document_path.write_text("MCP Tool 需要服务端授权。", encoding="utf-8")
            provider = Mock()
            provider.embed.side_effect = [[[0.9, 0.1]], [[1.0, 0.0]]]
            collection = Mock()
            collection.get.return_value = {"ids": []}
            collection.count.return_value = 1
            collection.query.return_value = {
                "documents": [["MCP Tool 需要服务端授权。"]],
                "metadatas": [[{"source": "mcp.md"}]],
                "distances": [[0.1]],
            }
            chroma_client = Mock()
            chroma_client.get_or_create_collection.return_value = collection

            knowledge = PersistentVectorKnowledgeBase(
                LocalKnowledgeBase(directory),
                provider,
                Path(directory) / "vectors",
                query_instruction="retrieve MCP documents",
                chroma_client=chroma_client,
            )
            chunks = knowledge.retrieve("MCP 授权", top_k=3)

            collection.upsert.assert_called_once()
            self.assertEqual(provider.embed.call_count, 2)
            self.assertEqual(
                provider.embed.call_args_list[0].kwargs["instruction"], ""
            )
            self.assertEqual(
                provider.embed.call_args_list[1].kwargs["instruction"],
                "retrieve MCP documents",
            )
            self.assertEqual(chunks[0].source, "mcp.md")
            self.assertEqual(chunks[0].score, 0.9)

    def test_sync_deletes_vectors_for_removed_documents(self) -> None:
        with TemporaryDirectory() as directory:
            collection = Mock()
            collection.get.return_value = {"ids": ["stale-id"]}
            chroma_client = Mock()
            chroma_client.get_or_create_collection.return_value = collection
            knowledge = PersistentVectorKnowledgeBase(
                LocalKnowledgeBase(directory),
                Mock(),
                Path(directory) / "vectors",
                chroma_client=chroma_client,
            )
            self.assertEqual(knowledge.sync(), 0)
            collection.delete.assert_called_once_with(ids=["stale-id"])

    def test_sync_updates_metadata_without_regenerating_embeddings(self) -> None:
        with TemporaryDirectory() as directory:
            document = Path(directory) / "tools.md"
            document.write_text(
                "# MCP Tools\n\nMCP Tools are model-controlled operations with schemas.",
                encoding="utf-8",
            )
            documents = LocalKnowledgeBase(directory)
            chunk = documents.chunks()[0]
            chunk_id = PersistentVectorKnowledgeBase._chunk_id(chunk)
            collection = Mock()
            collection.get.return_value = {
                "ids": [chunk_id],
                "metadatas": [{"source": "tools.md"}],
            }
            chroma_client = Mock()
            chroma_client.get_or_create_collection.return_value = collection
            provider = Mock()
            knowledge = PersistentVectorKnowledgeBase(
                documents,
                provider,
                Path(directory) / "vectors",
                chroma_client=chroma_client,
            )
            self.assertEqual(knowledge.sync(), 0)
            provider.embed.assert_not_called()
            collection.update.assert_called_once()
            metadata = collection.update.call_args.kwargs["metadatas"][0]
            self.assertEqual(metadata["title"], "MCP Tools")

    def test_embedding_provider_preserves_api_result_order(self) -> None:
        response = Mock()
        response.raise_for_status.return_value = None
        response.json.return_value = {
            "data": [
                {"index": 1, "embedding": [0.0, 1.0]},
                {"index": 0, "embedding": [1.0, 0.0]},
            ]
        }
        session = Mock()
        session.post.return_value = response
        provider = OpenAIEmbeddingProvider(
            "https://embedding.test/v1",
            "key",
            "embedding-model",
            dimensions=1024,
            failover_enabled=True,
            session=session,
        )
        self.assertEqual(
            provider.embed(["first", "second"], instruction="retrieve documents"),
            [[1.0, 0.0], [0.0, 1.0]],
        )
        request = session.post.call_args
        self.assertEqual(request.kwargs["headers"]["X-Failover-Enabled"], "true")
        self.assertEqual(request.kwargs["json"]["dimensions"], 1024)
        self.assertEqual(
            request.kwargs["json"]["instruction"], "retrieve documents"
        )


class KnowledgeImporterTests(unittest.TestCase):
    def test_html_removes_navigation_scripts_and_footer(self) -> None:
        html = """
        <html><head><title>MCP 安全</title><script>bad()</script></head>
        <body><nav>导航栏</nav><article><h1>MCP Tool 安全</h1>
        <p>MCP Server 必须把工具参数视为不可信输入，并在服务端进行授权检查。</p>
        </article><footer>版权信息</footer></body></html>
        """
        document = clean_html(html, "https://example.test/mcp-security")
        self.assertEqual(document.title, "MCP Tool 安全")
        self.assertIn("不可信输入", document.markdown)
        self.assertNotIn("导航栏", document.markdown)
        self.assertNotIn("版权信息", document.markdown)
        self.assertNotIn("bad()", document.markdown)

    def test_arxiv_page_keeps_only_title_and_abstract(self) -> None:
        html = """
        <html><body><nav>browse</nav><h1 class="title">Title: Retrieval-Augmented Generation</h1>
        <blockquote class="abstract">Abstract: A sufficiently long description of retrieval-augmented
        generation that combines parametric and non-parametric memory for knowledge-intensive tasks.</blockquote>
        <section>Submission history and unrelated recommendation widgets</section></body></html>
        """
        document = clean_html(html, "https://arxiv.org/abs/2005.11401")
        self.assertEqual(document.title, "Retrieval-Augmented Generation")
        self.assertIn("non-parametric memory", document.markdown)
        self.assertNotIn("Submission history", document.markdown)

    def test_docx_preserves_headings_and_tables(self) -> None:
        from docx import Document

        document = Document()
        document.add_heading("MCP 实践指南", level=1)
        document.add_paragraph("这是一段足够长的 MCP 知识库说明，用于验证 Word 正文提取与格式保留。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "能力"
        table.cell(0, 1).text = "用途"
        table.cell(1, 0).text = "Tools"
        table.cell(1, 1).text = "执行操作"
        data = BytesIO()
        document.save(data)

        cleaned = clean_docx(data.getvalue(), "guide.docx")
        self.assertEqual(cleaned.title, "MCP 实践指南")
        self.assertIn("# MCP 实践指南", cleaned.markdown)
        self.assertIn("| 能力 | 用途 |", cleaned.markdown)

    def test_scanned_pdf_is_rejected_with_ocr_message(self) -> None:
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        data = BytesIO()
        writer.write(data)
        with self.assertRaisesRegex(ValueError, "OCR"):
            clean_pdf(data.getvalue(), "scan.pdf")

    def test_imports_local_markdown_to_imported_directory(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.md"
            source.write_text(
                "# MCP Resources\n\nMCP Resources 用于向客户端提供可读取的上下文，"
                "应保留稳定的 URI 和清晰的来源信息。",
                encoding="utf-8",
            )
            output = root / "knowledge" / "imported"
            result = KnowledgeImporter(output).import_one(str(source))
            saved = result.output_path.read_text(encoding="utf-8")
            self.assertTrue(result.output_path.is_file())
            self.assertEqual(saved.count("# MCP Resources"), 1)
            self.assertEqual(len(re.findall(r"^#\s+", saved, flags=re.M)), 1)
            self.assertIn("<!-- source:", saved)

    def test_url_markdown_is_not_treated_as_html(self) -> None:
        response = Mock()
        response.raise_for_status.return_value = None
        response.headers = {"Content-Type": "text/plain; charset=utf-8"}
        response.encoding = "utf-8"
        response.iter_content.return_value = [
            "# MCP Transport\n\nstdio 传输适合本地子进程通信，"
            "Streamable HTTP 适合远程服务连接。".encode("utf-8")
        ]
        session = Mock()
        session.get.return_value = response
        with TemporaryDirectory() as directory:
            result = KnowledgeImporter(directory, session=session).import_one(
                "https://raw.example.test/transport.md"
            )
            saved = result.output_path.read_text(encoding="utf-8")
            self.assertIn("stdio 传输", saved)
            self.assertNotIn("<html", saved)

    def test_markdown_finds_title_after_badges(self) -> None:
        response = Mock()
        response.raise_for_status.return_value = None
        response.headers = {"Content-Type": "text/plain"}
        response.encoding = "utf-8"
        response.iter_content.return_value = [
            b"[build badge](https://example.test)\n\n# Chroma\n\n"
            b"Chroma is an open-source database for AI applications with embeddings."
        ]
        session = Mock()
        session.get.return_value = response
        with TemporaryDirectory() as directory:
            result = KnowledgeImporter(directory, session=session).import_one(
                "https://raw.example.test/README.md"
            )
            self.assertEqual(result.title, "Chroma")

    def test_markdown_code_comments_are_not_treated_as_headings(self) -> None:
        response = Mock()
        response.raise_for_status.return_value = None
        response.headers = {"Content-Type": "text/plain"}
        response.encoding = "utf-8"
        response.iter_content.return_value = [
            b"```bash\npip install chromadb\n# this is a shell comment\n```\n\n"
            b"A sufficiently long explanation about persistent vector storage and retrieval."
        ]
        session = Mock()
        session.get.return_value = response
        with TemporaryDirectory() as directory:
            result = KnowledgeImporter(directory, session=session).import_one(
                "https://raw.githubusercontent.com/chroma-core/chroma/main/README.md"
            )
            saved = result.output_path.read_text(encoding="utf-8")
            self.assertEqual(result.title, "chroma")
            self.assertIn("# this is a shell comment", saved)
            self.assertNotIn("## this is a shell comment", saved)

    def test_mdx_removes_frontmatter_imports_and_component_tags(self) -> None:
        response = Mock()
        response.raise_for_status.return_value = None
        response.headers = {"Content-Type": "text/plain"}
        response.encoding = "utf-8"
        response.iter_content.return_value = [
            b"---\ntitle: MCP Tool Specification\n---\nimport Card from './Card'\n"
            b"<Note>Tools are model-controlled capabilities with schemas and execution rules.</Note>"
        ]
        session = Mock()
        session.get.return_value = response
        with TemporaryDirectory() as directory:
            result = KnowledgeImporter(directory, session=session).import_one(
                "https://raw.example.test/tools.mdx"
            )
            saved = result.output_path.read_text(encoding="utf-8")
            self.assertIn("# MCP Tool Specification", saved)
            self.assertIn("Tools are model-controlled", saved)
            self.assertNotIn("import Card", saved)
            self.assertNotIn("<Note>", saved)

    def test_reimport_replaces_old_file_for_same_source(self) -> None:
        with TemporaryDirectory() as directory:
            source = Path(directory) / "source.md"
            source.write_text(
                "# First Title\n\n第一版内容足够长，用于验证同一来源重新导入时会替换旧文件。"
                "这些补充文字用来模拟真实的知识文档段落。",
                encoding="utf-8",
            )
            output = Path(directory) / "imported"
            importer = KnowledgeImporter(output)
            first = importer.import_one(str(source))
            source.write_text(
                "# Second Title\n\n第二版内容足够长，用于验证内容和标题变更后不会留下重复文档。"
                "这些补充文字用来模拟更新后的知识文档段落。",
                encoding="utf-8",
            )
            second = importer.import_one(str(source))
            self.assertFalse(first.output_path.exists())
            self.assertTrue(second.output_path.exists())
            self.assertEqual(len(list(output.glob("*.md"))), 1)


class SchedulerTests(unittest.TestCase):
    def test_first_run_is_due_and_then_waits_twelve_hours(self) -> None:
        now = datetime(2026, 7, 30, 9, 0, tzinfo=timezone.utc)
        self.assertTrue(_is_due(now, None, 12))
        self.assertFalse(_is_due(now + timedelta(hours=11, minutes=59), now, 12))
        self.assertTrue(_is_due(now + timedelta(hours=12), now, 12))

    def test_persists_last_success(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "scheduler.json"
            published_at = datetime(2026, 7, 30, 9, 0, tzinfo=timezone.utc)
            _mark_success(published_at, path)
            self.assertEqual(_last_success(path), published_at)


class CSDNClientTests(unittest.TestCase):
    def test_payload_contains_markdown_and_html(self) -> None:
        client = CSDNClient(settings())
        payload = client.build_payload(
            GeneratedArticle(
                title="测试",
                markdown="# 测试\n\n```python\nprint('ok')\n```",
                description="描述。",
                tags="Spring Boot 3,Redis",
            )
        )
        self.assertEqual(payload["markdowncontent"], "# 测试\n\n```python\nprint('ok')\n```")
        self.assertIn("<code", payload["content"])
        self.assertEqual(payload["tags"], "Spring Boot 3,Redis")
        self.assertEqual(payload["status"], 0)

    def test_payload_uses_configured_tags_as_fallback(self) -> None:
        client = CSDNClient(settings())
        payload = client.build_payload(
            GeneratedArticle(title="测试", markdown="# 测试\n\n正文", description="描述。")
        )
        self.assertEqual(payload["tags"], "python")

    def test_publish_checks_business_code(self) -> None:
        session = Mock()
        response = Mock()
        response.raise_for_status.return_value = None
        response.json.return_value = {"code": 401, "msg": "未登录"}
        session.post.return_value = response
        client = CSDNClient(settings(), session=session)
        with self.assertRaisesRegex(RuntimeError, "发布失败"):
            client.publish(
                GeneratedArticle(title="标题", markdown="# 标题\n\n正文", description="正文。")
            )


if __name__ == "__main__":
    unittest.main()
