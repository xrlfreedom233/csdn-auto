from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urlparse

import requests
from bs4 import BeautifulSoup
from markdownify import markdownify
from pypdf import PdfReader


SUPPORTED_LOCAL_SUFFIXES = {
    ".docx",
    ".htm",
    ".html",
    ".markdown",
    ".md",
    ".mdx",
    ".pdf",
    ".txt",
}


@dataclass(frozen=True)
class ImportedDocument:
    title: str
    markdown: str
    source: str


@dataclass(frozen=True)
class ImportResult:
    source: str
    output_path: Path
    title: str
    characters: int


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _safe_title(value: str, fallback: str = "imported-document") -> str:
    title = re.sub(r"\s+", " ", value).strip().strip("#")
    return title[:160] or fallback


def _slug(value: str) -> str:
    value = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff_-]+", "-", value)
    return value.strip("-")[:70] or "document"


def _first_markdown_h1(text: str) -> str:
    in_fence = False
    for line in text.splitlines():
        if re.match(r"^\s*(?:```|~~~)", line):
            in_fence = not in_fence
            continue
        if not in_fence:
            match = re.match(r"^#\s+(.+?)\s*$", line)
            if match:
                return match.group(1)
    return ""


def _normalize_markdown_body(text: str) -> str:
    """Remove the document H1 and demote later H1s without touching code."""
    in_fence = False
    removed_title = False
    lines: list[str] = []
    for line in text.splitlines():
        if re.match(r"^\s*(?:```|~~~)", line):
            in_fence = not in_fence
            lines.append(line)
            continue
        if not in_fence and re.match(r"^#\s+", line):
            if not removed_title:
                removed_title = True
                continue
            line = "#" + line
        lines.append(line)
    return _clean_text("\n".join(lines))


def _document_from_text(text: str, source: str) -> ImportedDocument:
    metadata_title = ""
    if Path(urlparse(source).path).suffix.lower() == ".mdx":
        frontmatter = re.match(r"\A---\s*\n(.*?)\n---\s*\n", text, flags=re.S)
        if frontmatter:
            title_match = re.search(
                r"^title:\s*['\"]?(.+?)['\"]?\s*$",
                frontmatter.group(1),
                flags=re.M,
            )
            if title_match:
                metadata_title = title_match.group(1).strip()
        text = re.sub(r"\A---\s*\n.*?\n---\s*\n", "", text, flags=re.S)
        text = re.sub(r"^(?:import|export)\s+.+$", "", text, flags=re.M)
        text = re.sub(r"\{\/\*.*?\*\/\}", "", text, flags=re.S)
        text = re.sub(r"<\/?[A-Z][^>]*>", "", text)
    content = _clean_text(text)
    heading = _first_markdown_h1(content)
    title = _safe_title(
        heading or metadata_title,
        _title_from_source(source),
    )
    return ImportedDocument(title, content, source)


def clean_html(html: str, source: str) -> ImportedDocument:
    soup = BeautifulSoup(html, "html.parser")
    page_title = soup.find("title")
    for node in soup.select(
        "script, style, noscript, template, svg, canvas, nav, footer, aside, "
        "form, dialog, iframe, [aria-hidden='true']"
    ):
        node.decompose()
    if urlparse(source).netloc.lower().endswith("arxiv.org"):
        paper_title = soup.select_one("h1.title")
        abstract = soup.select_one("blockquote.abstract")
        if paper_title and abstract:
            title = re.sub(
                r"^\s*Title:\s*", "", paper_title.get_text(" ", strip=True)
            )
            content = re.sub(
                r"^\s*Abstract:\s*", "", abstract.get_text(" ", strip=True)
            )
            return ImportedDocument(_safe_title(title), _clean_text(content), source)
    root = soup.find("article") or soup.find("main") or soup.body or soup
    title_node = root.find("h1") or page_title
    title = _safe_title(
        title_node.get_text(" ", strip=True) if title_node else "",
        _title_from_source(source),
    )
    for node in root.select(
        ".advertisement, .ads, .breadcrumb, .comment, .comments, .cookie, "
        ".footer, .header, .menu, .nav, .pagination, .recommend, .related, "
        ".share, .sidebar, [class*='advert'], [id*='advert']"
    ):
        node.decompose()
    content = markdownify(str(root), heading_style="ATX", bullets="-")
    content = re.sub(r"\n\s*([*-])\s*\n", "\n", content)
    return ImportedDocument(title, _clean_text(content), source)


def clean_pdf(data: bytes, source: str) -> ImportedDocument:
    reader = PdfReader(BytesIO(data))
    pages = [_clean_text(page.extract_text() or "") for page in reader.pages]
    nonempty_pages = [page for page in pages if page]
    if not nonempty_pages:
        raise ValueError("未从 PDF 提取到文字，该文件可能是扫描版，需先进行 OCR")

    line_pages = [[line.strip() for line in page.splitlines() if line.strip()] for page in pages]
    repeated_edges: set[str] = set()
    if len(nonempty_pages) >= 3:
        edge_counts = Counter(
            line
            for lines in line_pages
            for line in set(lines[:1] + lines[-1:])
            if len(line) <= 120
        )
        threshold = max(3, (len(nonempty_pages) * 3 + 4) // 5)
        repeated_edges = {
            line for line, count in edge_counts.items() if count >= threshold
        }
    cleaned_pages: list[str] = []
    for lines in line_pages:
        lines = [line for line in lines if line not in repeated_edges]
        text = "\n".join(lines)
        text = re.sub(r"(?<=[\u4e00-\u9fff，。；：])\n(?=[\u4e00-\u9fff])", "", text)
        text = re.sub(r"(?<=[A-Za-z,;:])\n(?=[a-z])", " ", text)
        if text.strip():
            cleaned_pages.append(text.strip())
    title = _title_from_source(source)
    metadata_title = reader.metadata.title if reader.metadata else None
    if metadata_title:
        title = _safe_title(str(metadata_title), title)
    return ImportedDocument(title, _clean_text("\n\n".join(cleaned_pages)), source)


def clean_docx(data: bytes, source: str) -> ImportedDocument:
    from docx import Document

    document = Document(BytesIO(data))
    parts: list[str] = []
    title = ""
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "").lower()
        heading_match = re.search(r"heading\s*(\d+)|标题\s*(\d+)", style_name)
        if style_name in {"title", "主标题"}:
            title = title or text
            parts.append(f"# {text}")
        elif heading_match:
            level = int(heading_match.group(1) or heading_match.group(2) or 2)
            level = min(max(level, 1), 6)
            title = title or (text if level == 1 else "")
            parts.append(f"{'#' * level} {text}")
        elif style_name.startswith("list"):
            parts.append(f"- {text}")
        else:
            parts.append(text)

    for table in document.tables:
        rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        parts.append("| " + " | ".join(normalized[0]) + " |")
        parts.append("| " + " | ".join(["---"] * width) + " |")
        parts.extend("| " + " | ".join(row) + " |" for row in normalized[1:])

    content = _clean_text("\n\n".join(parts))
    if not content:
        raise ValueError("Word 文档中没有可导入的文字")
    return ImportedDocument(_safe_title(title, _title_from_source(source)), content, source)


def _title_from_source(source: str) -> str:
    parsed = urlparse(source)
    if parsed.scheme in {"http", "https"}:
        parts = [item for item in unquote(parsed.path).split("/") if item]
        if (
            parsed.netloc == "raw.githubusercontent.com"
            and len(parts) >= 4
            and parts[-1].lower().startswith("readme")
        ):
            candidate = parts[1].replace("-", " ")
        else:
            candidate = Path(unquote(parsed.path)).stem or parsed.netloc
    else:
        candidate = Path(source).stem
    return _safe_title(candidate)


class KnowledgeImporter:
    def __init__(
        self,
        output_directory: str | Path = "knowledge/imported",
        timeout: int = 30,
        max_bytes: int = 25 * 1024 * 1024,
        session: requests.Session | None = None,
    ):
        self.output_directory = Path(output_directory)
        self.timeout = timeout
        self.max_bytes = max_bytes
        self.session = session or requests.Session()

    def import_many(self, sources: list[str]) -> list[ImportResult]:
        expanded: list[str] = []
        for source in sources:
            path = Path(source)
            if path.is_dir():
                expanded.extend(
                    str(item)
                    for item in sorted(path.rglob("*"))
                    if item.is_file() and item.suffix.lower() in SUPPORTED_LOCAL_SUFFIXES
                )
            else:
                expanded.append(source)
        return [self.import_one(source) for source in expanded]

    def import_one(self, source: str) -> ImportResult:
        document = self._read_url(source) if _is_url(source) else self._read_file(source)
        body = _normalize_markdown_body(document.markdown)
        if len(body) < 40:
            raise ValueError(f"清洗后正文过短，拒绝导入：{source}")
        digest = sha256(body.encode("utf-8")).hexdigest()[:12]
        output_path = self.output_directory / f"{_slug(document.title)}-{digest}.md"
        self.output_directory.mkdir(parents=True, exist_ok=True)
        imported_at = datetime.now().astimezone().isoformat()
        safe_source = document.source.replace("--", "—")
        source_header = f"<!-- source: {safe_source} -->"
        for existing in self.output_directory.glob("*.md"):
            if existing == output_path:
                continue
            with existing.open(encoding="utf-8") as stream:
                first_line = stream.readline().rstrip("\n")
            if first_line == source_header:
                existing.unlink()
        markdown = (
            f"{source_header}\n"
            f"<!-- imported_at: {imported_at} -->\n\n"
            f"# {document.title}\n\n{body}\n"
        )
        output_path.write_text(markdown, encoding="utf-8")
        return ImportResult(source, output_path, document.title, len(body))

    def _read_url(self, source: str) -> ImportedDocument:
        try:
            response = self.session.get(
                source,
                headers={"User-Agent": "csdn-auto-knowledge-importer/1.0"},
                timeout=self.timeout,
                stream=True,
            )
            response.raise_for_status()
            data = self._limited_content(response)
        except requests.RequestException as exc:
            raise RuntimeError(f"下载文档失败：{source}：{exc}") from exc
        content_type = response.headers.get("Content-Type", "").lower()
        suffix = Path(urlparse(source).path).suffix.lower()
        if "pdf" in content_type or suffix == ".pdf":
            return clean_pdf(data, source)
        if "wordprocessingml" in content_type or suffix == ".docx":
            return clean_docx(data, source)
        encoding = response.encoding or "utf-8"
        text = data.decode(encoding, errors="replace")
        if suffix in {".md", ".markdown", ".mdx", ".txt"} or any(
            item in content_type
            for item in ("text/markdown", "text/plain", "text/x-markdown")
        ):
            return _document_from_text(text, source)
        return clean_html(text, source)

    def _limited_content(self, response: requests.Response) -> bytes:
        length = response.headers.get("Content-Length")
        if length and int(length) > self.max_bytes:
            raise ValueError(f"文档超过 {self.max_bytes // 1024 // 1024} MB 限制")
        chunks: list[bytes] = []
        total = 0
        for chunk in response.iter_content(chunk_size=64 * 1024):
            total += len(chunk)
            if total > self.max_bytes:
                raise ValueError(f"文档超过 {self.max_bytes // 1024 // 1024} MB 限制")
            chunks.append(chunk)
        return b"".join(chunks)

    def _read_file(self, source: str) -> ImportedDocument:
        path = Path(source)
        if not path.is_file():
            raise FileNotFoundError(f"文档不存在：{source}")
        if path.stat().st_size > self.max_bytes:
            raise ValueError(f"文档超过 {self.max_bytes // 1024 // 1024} MB 限制：{source}")
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_LOCAL_SUFFIXES:
            raise ValueError(f"不支持的文档格式：{suffix or '无扩展名'}")
        data = path.read_bytes()
        if suffix == ".pdf":
            return clean_pdf(data, str(path))
        if suffix == ".docx":
            return clean_docx(data, str(path))
        text = data.decode("utf-8")
        if suffix in {".html", ".htm"}:
            return clean_html(text, str(path))
        return _document_from_text(text, str(path))


def _is_url(source: str) -> bool:
    return urlparse(source).scheme.lower() in {"http", "https"}
